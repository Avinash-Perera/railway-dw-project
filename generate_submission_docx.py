"""
Generates the assignment submission document (DOCX + Markdown spec) from the live data warehouse.
KPI actual values are queried from the database so the document always matches the loaded data.

Usage:  MYSQL_PORT=3306 python3 generate_submission_docx.py
"""
import os
import sys
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
from etl.db_connector import get_engine

FIG = os.path.join(ROOT, "paper", "figures")
OUT_DOCX = os.path.join(ROOT, "Railway_DW_Design_Submission.docx")
OUT_MD = os.path.join(ROOT, "KPI_AND_DATA_MODEL_SPECIFICATION.md")

# ---------------------------------------------------------------- live data
def q(engine, sql):
    return pd.read_sql(sql, engine)

def collect(engine):
    d = {}
    d['counts'] = q(engine, """
        SELECT 'Fact_TrainDelay' t, COUNT(*) n FROM Fact_TrainDelay UNION ALL
        SELECT 'Fact_TrainSchedule', COUNT(*) FROM Fact_TrainSchedule UNION ALL
        SELECT 'Fact_JourneyDeparture', COUNT(*) FROM Fact_JourneyDeparture UNION ALL
        SELECT 'Fact_JourneyArrival', COUNT(*) FROM Fact_JourneyArrival UNION ALL
        SELECT 'Fact_JourneyCancellation', COUNT(*) FROM Fact_JourneyCancellation UNION ALL
        SELECT 'Fact_DelayIncident', COUNT(*) FROM Fact_DelayIncident UNION ALL
        SELECT 'Fact_WeatherObservation', COUNT(*) FROM Fact_WeatherObservation UNION ALL
        SELECT 'Fact_TrainDailySnapshot', COUNT(*) FROM Fact_TrainDailySnapshot UNION ALL
        SELECT 'Dim_Date', COUNT(*) FROM Dim_Date UNION ALL
        SELECT 'Dim_Time', COUNT(*) FROM Dim_Time UNION ALL
        SELECT 'Dim_Station', COUNT(*) FROM Dim_Station UNION ALL
        SELECT 'Dim_Train', COUNT(*) FROM Dim_Train UNION ALL
        SELECT 'Dim_DelayCategory', COUNT(*) FROM Dim_DelayCategory UNION ALL
        SELECT 'Dim_Weather', COUNT(*) FROM Dim_Weather UNION ALL
        SELECT 'Dim_DelayReason', COUNT(*) FROM Dim_DelayReason UNION ALL
        SELECT 'Dim_CancellationReason', COUNT(*) FROM Dim_CancellationReason""")
    d['span'] = q(engine, "SELECT MIN(d.FullDate) a, MAX(d.FullDate) b FROM Fact_TrainDelay f JOIN Dim_Date d ON f.DateSK=d.DateSK").iloc[0]
    d['overall'] = q(engine, """
        SELECT SUM(PunctualFlag)/SUM(TripCount)*100 k1,
               SUM(SevereDelayFlag)/SUM(TripCount)*100 k2,
               SUM(ExcessDelayMinutes)/SUM(TripCount) k3
        FROM Fact_TrainDelay""").iloc[0]
    d['k4'] = q(engine, """
        SELECT ((SUM(CASE WHEN d.IsWeekend=1 THEN f.DelayMinutes END)/SUM(CASE WHEN d.IsWeekend=1 THEN f.TripCount END)) /
                (SUM(CASE WHEN d.IsWeekend=0 THEN f.DelayMinutes END)/SUM(CASE WHEN d.IsWeekend=0 THEN f.TripCount END)) - 1)*100 v
        FROM Fact_TrainDelay f JOIN Dim_Date d ON f.DateSK=d.DateSK""").iloc[0]['v']
    d['k5'] = q(engine, """
        SELECT ((SUM(CASE WHEN t.IsPeakHour=1 THEN f.DelayMinutes END)/SUM(CASE WHEN t.IsPeakHour=1 THEN f.TripCount END)) /
                (SUM(CASE WHEN t.IsPeakHour=0 THEN f.DelayMinutes END)/SUM(CASE WHEN t.IsPeakHour=0 THEN f.TripCount END)) - 1)*100 v
        FROM Fact_TrainDelay f JOIN Dim_Time t ON f.TimeSK=t.TimeSK WHERE t.TimeSK<>9999""").iloc[0]['v']
    d['monthly'] = q(engine, """
        SELECT m.Year, m.Month, m.TotalTrips, m.PunctualityRatePct, m.PunctualityStatus, m.PunctualityMoMChangePts,
               m.SevereDelayRatePct, m.SevereDelayStatus, m.SevereDelayMoMChangePts,
               m.AvgExcessDelayMin, m.ExcessDelayStatus, m.ExcessDelayMoMChangeMin,
               w.WeekendUpliftPct, w.WeekendUpliftStatus, w.WeekendUpliftMoMChangePts,
               p.PeakSurgePct, p.PeakSurgeStatus, p.PeakSurgeMoMChangePts
        FROM vw_KPI_Monthly m
        JOIN vw_KPI_WeekendUplift w ON w.Year=m.Year AND w.Month=m.Month
        JOIN vw_KPI_PeakSurge p ON p.Year=m.Year AND p.Month=m.Month
        ORDER BY m.Year, m.Month""")
    return d

def status(v, green, amber, lower_is_better=True):
    if lower_is_better:
        return 'Green' if v <= green else ('Amber' if v <= amber else 'Red')
    return 'Green' if v >= green else ('Amber' if v >= amber else 'Red')

# ---------------------------------------------------------------- doc helpers
doc = Document()
md = []

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def H(text, level=1):
    doc.add_heading(text, level=level)
    md.append(f"\n{'#' * (level + 1)} {text}\n")

def P(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    md.append(('**%s**' % text if bold else text) + "\n")
    return p

def B(items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')
        md.append(f"- {it}")
    md.append("")

def CODE(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = 'Consolas'; r.font.size = Pt(8)
    md.append("```sql\n" + text + "\n```\n")

def T(headers, rows, widths=None, font=8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(font)
        run.font.color.rgb = RGBColor(255, 255, 255); shade(c, '1F3A5F')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(font)
            if str(v) in ('Green', 'Amber', 'Red'):
                shade(cells[i], {'Green': 'C6EFCE', 'Amber': 'FFEB9C', 'Red': 'FFC7CE'}[str(v)])
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    md.append("| " + " | ".join(headers) + " |")
    md.append("|" + "|".join(["---"] * len(headers)) + "|")
    for row in rows:
        md.append("| " + " | ".join(str(v).replace("|", "/") for v in row) + " |")
    md.append("")

def IMG(path, caption, width=6.5):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = doc.add_paragraph(caption); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic = True
    md.append(f"![{caption}](paper/figures/{os.path.basename(path)})\n")


# ---------------------------------------------------------------- content
def build(d):
    n = {r.t: int(r.n) for r in d['counts'].itertuples()}
    span_a, span_b = d['span']['a'], d['span']['b']
    o = d['overall']
    k1, k2, k3, k4, k5 = float(o.k1), float(o.k2), float(o.k3), float(d['k4']), float(d['k5'])
    s1, s2, s3 = status(k1, 75, 60, False), status(k2, 10, 15), status(k3, 10, 20)
    s4, s5 = status(k4, 5, 15), status(k5, 10, 25)
    mon = d['monthly']
    mon = mon[mon['TotalTrips'] > 1_000_000] if (mon['TotalTrips'] > 1_000_000).any() else mon  # complete months only

    title = doc.add_heading('Indian Railways Data Warehouse', 0)
    md.append("# Indian Railways Data Warehouse\n## Dimensional Model, Business Events & KPI Specification\n")
    doc.add_heading('Dimensional Model, Business Events & KPI Specification', 1)
    P("Course deliverable: Data Warehouse Design (Dimensional Model) with ETL implementation as bonus.", italic=True)

    H("1. Project Overview & Data Sources")
    P("The warehouse analyses train punctuality across the Indian Railways network. It is modelled as a fact "
      "constellation (several star schemas sharing conformed dimensions) in MySQL 8, following Kimball dimensional modelling.")
    T(["Source file", "Rows", "Used for"], [
        ["combined_delay.csv", "38,428,703", "Fact_TrainDelay (delay per train/station/date), Fact_JourneyDeparture (stop 1), Fact_JourneyArrival (last stop)"],
        ["combined_schedule.csv", "172,112", "Fact_TrainSchedule (timetable distance/dwell); stop sequence -> Dim_Time scheduled lookup"],
        ["train_details.csv", "8,992", "Dim_Train (SCD Type 2 with start/end dates and active flag)"],
        ["station_full_names.csv", "8,963", "Dim_Station (spatial dimension with zone hierarchy)"],
        ["train_delays.csv", "30", "Fact_JourneyCancellation, Fact_DelayIncident, Fact_WeatherObservation"],
    ], [2.0, 1.0, 3.5])
    P(f"Observation window in the delay data: {span_a} to {span_b}.")

    H("2. Business Events (8 identified and modeled)")
    P("Each business event is represented by an independent fact table. All 8 events are fully modelled (DDL, grain, keys, measures) "
      "and physically loaded with verified data in the MySQL 8 data warehouse.")
    T(["#", "Business event", "Fact table", "Grain (one row =)", "Key measures", "Status"], [
        ["1", "Train arrival at a station", "Fact_TrainDelay", "One train arriving at one station on one date",
         "DelayMinutes, TripCount, IsDelayed, PunctualFlag, SevereDelayFlag, ExcessDelayMinutes", f"Loaded ({n['Fact_TrainDelay']:,} rows)"],
        ["2", "Scheduled stop of a train", "Fact_TrainSchedule", "One scheduled stop of one train at one station",
         "DistanceFromOriginKm, StopCount, ScheduledDwellMinutes, SegmentDistanceKm", f"Loaded ({n['Fact_TrainSchedule']:,} rows)"],
        ["3", "Journey departure", "Fact_JourneyDeparture", "One journey leaving its origin on one date",
         "DepartureDelayMinutes, DepartureCount, DelayedDepartureFlag", f"Loaded ({n['Fact_JourneyDeparture']:,} rows)"],
        ["4", "Journey arrival", "Fact_JourneyArrival", "One journey reaching its destination on one date",
         "ArrivalDelayMinutes, ArrivalCount, DelayedArrivalFlag", f"Loaded ({n['Fact_JourneyArrival']:,} rows)"],
        ["5", "Journey cancellation", "Fact_JourneyCancellation", "One cancelled journey on one date",
         "CancellationCount", f"Loaded ({n['Fact_JourneyCancellation']:,} rows)"],
        ["6", "Delay incident (cause)", "Fact_DelayIncident", "One delayed journey attributed to one primary reason",
         "IncidentDelayMinutes, IncidentCount", f"Loaded ({n['Fact_DelayIncident']:,} rows)"],
        ["7", "Weather observation", "Fact_WeatherObservation", "One weather condition for one journey on one date",
         "ObservationCount, JourneyDelayMinutes, WeatherAffectedFlag", f"Loaded ({n['Fact_WeatherObservation']:,} rows)"],
        ["8", "Daily train snapshot", "Fact_TrainDailySnapshot", "One train on one date (periodic snapshot)",
         "StopsRecorded, DelayedStops, MaxDelayMinutes, TotalDelayMinutes, TotalExcessDelayMinutes", f"Loaded ({n['Fact_TrainDailySnapshot']:,} rows)"],
    ], [0.3, 1.2, 1.4, 1.6, 1.8, 0.9], font=7)

    H("3. Enterprise Bus Matrix")
    IMG(os.path.join(FIG, "fig_bus_matrix.png"), "Figure 1 - Bus matrix: 8 events x 8 conformed dimensions", 6.8)

    H("4. Star Schema Model")
    IMG(os.path.join(FIG, "fig_star_traindelay.png"), "Figure 2 - Star schema of Fact_TrainDelay with explicit foreign-key join lines", 6.6)
    H("4.1 Dimension tables", 2)
    T(["Dimension", "Surrogate PK", "Natural key", "Attributes", "Hierarchy (drill-down order)"], [
        ["Dim_Date", "DateSK (YYYYMMDD)", "FullDate", "Year, Quarter, QuarterName, Month, MonthName, Day, WeekOfYear, DayOfWeek, IsWeekend, IsHoliday, HolidayName, Season", "Year > Quarter > Month > Day"],
        ["Dim_Time", "TimeSK (HHMM; 9999 = Unknown)", "TimeLabel", "Hour24, Minute, TimeOfDayBand, IsPeakHour", "TimeOfDayBand > Hour24 > Minute"],
        ["Dim_Station", "StationSK", "StationCode", "StationName, ZoneCode, ZoneName", "ZoneName > StationName"],
        ["Dim_Train (SCD2)", "TrainSK", "TrainNumber", "TrainName, TrainTypeCode, TrainTypeName, ServiceClass, EffectiveStartDate, EffectiveEndDate, IsCurrent", "ServiceClass > TrainTypeName > TrainNumber"],
        ["Dim_DelayCategory", "DelayCatSK", "-", "CategoryName, MinDelayMinutes, MaxDelayMinutes", "-"],
        ["Dim_Weather", "WeatherSK", "-", "WeatherCondition, WeatherSeverity", "WeatherSeverity > WeatherCondition"],
        ["Dim_DelayReason", "DelayReasonSK", "-", "DelayReasonName, DelayReasonGroup", "DelayReasonGroup > DelayReasonName"],
        ["Dim_CancellationReason", "CancellationReasonSK", "-", "CancellationReasonName, CancellationGroup", "CancellationGroup > CancellationReasonName"],
    ], [1.2, 1.2, 0.8, 2.4, 1.6], font=7)
    P(f"Loaded dimension rows: Dim_Date {n['Dim_Date']:,}; Dim_Time {n['Dim_Time']:,}; Dim_Station {n['Dim_Station']:,}; "
      f"Dim_Train {n['Dim_Train']:,}; Dim_DelayCategory {n['Dim_DelayCategory']:,}; Dim_Weather {n['Dim_Weather']:,}; "
      f"Dim_DelayReason {n['Dim_DelayReason']:,}; Dim_CancellationReason {n['Dim_CancellationReason']:,}.")

    H("4.2 Measures of Fact_TrainDelay", 2)
    T(["Measure", "Type", "Derivation", "Additivity"], [
        ["DelayMinutes", "Raw", "Source column delay (negative = early)", "Additive"],
        ["TripCount", "Derived", "Constant 1 per arrival event", "Additive"],
        ["IsDelayed", "Derived flag", "1 if DelayMinutes > 0", "Additive (count of delayed)"],
        ["PunctualFlag", "Derived flag", "1 if DelayMinutes <= 15 (punctuality tolerance)", "Additive"],
        ["SevereDelayFlag", "Derived flag", "1 if DelayMinutes > 45", "Additive"],
        ["ExcessDelayMinutes", "Derived", "GREATEST(0, DelayMinutes - 15)", "Additive"],
    ], [1.5, 1.0, 3.0, 1.5], font=8)

    H("5. Design Decisions vs. Marking Criteria")
    T(["Criterion", "How it is satisfied in this model"], [
        ["Derived measures", "Fact_TrainDelay stores calculated TripCount, IsDelayed, PunctualFlag, SevereDelayFlag, ExcessDelayMinutes in addition to raw DelayMinutes. Fact_TrainSchedule stores derived ScheduledDwellMinutes and SegmentDistanceKm."],
        ["No unnecessary dimension attributes", "station_address from the source file is deliberately NOT loaded into Dim_Station. Dim_Train carries only analytical attributes."],
        ["Surrogate vs natural keys", "Every dimension has a surrogate integer PK (StationSK, TrainSK, ...) and, where applicable, the natural key (StationCode, TrainNumber). Fact tables contain surrogate keys only; no operational keys."],
        ["SCD Type 2", "Dim_Train has EffectiveStartDate, EffectiveEndDate (open row = 9999-12-31) and IsCurrent, plus UNIQUE(TrainNumber, EffectiveStartDate). etl/load_dimensions.py::apply_train_scd2_change expires the current row and inserts a new version."],
        ["Date dimension enrichment", "Dim_Date includes IsWeekend, IsHoliday, HolidayName (Indian public holidays), Season, WeekOfYear, QuarterName."],
        ["Hierarchy consistency & order", "Year > Quarter > Month > Day; ZoneName > StationName; ServiceClass > TrainTypeName > TrainNumber. All levels live in ONE dimension (no split hierarchies) and all attributes physically exist."],
        ["Table relationships", "All fact tables declare FOREIGN KEY constraints to their dimensions; Figure 2 shows the join lines."],
        ["Naming conventions", "Fact_ / Dim_ prefixes, <Name>SK for surrogate keys, PascalCase everywhere; no mixed conventions."],
        ["Readable descriptors", "Source codes (e.g. EXP-TRAINS, NR) are expanded to Express / Northern Railway to avoid cryptic abbreviations."],
    ], [1.8, 5.0], font=8)
    P("SCD Type 2 example (train re-classified from Express to Superfast on 2026-04-01):", bold=True)
    CODE("UPDATE Dim_Train SET EffectiveEndDate = '2026-03-31', IsCurrent = 0\n"
         " WHERE TrainNumber = '12951' AND IsCurrent = 1;\n"
         "INSERT INTO Dim_Train (TrainNumber, TrainName, TrainTypeCode, TrainTypeName, ServiceClass,\n"
         "                       EffectiveStartDate, EffectiveEndDate, IsCurrent)\n"
         "VALUES ('12951', 'Mumbai Rajdhani', 'SF-TRAINS', 'Superfast', 'Mail / Express', '2026-04-01', '9999-12-31', 1);")

    H("6. Key Performance Indicators (5)")
    P("Every KPI has the four mandatory properties: Value (ratio/percentage formula), a single numeric Goal, "
      "a Status band (Green/Amber/Red) and a Trend (comparison with the prior period). Actual values are computed "
      "from the loaded warehouse by the views vw_KPI_Monthly, vw_KPI_WeekendUplift and vw_KPI_PeakSurge.")
    kpis = [
        ("KPI 1 - Punctuality Rate (%)", "SUM(PunctualFlag) / SUM(TripCount) x 100", ">= 75%",
         "Green >= 75% | Amber 60% to < 75% | Red < 60%", "Month-over-Month change (percentage points)", f"{k1:.2f}%", s1, "PunctualFlag, TripCount"),
        ("KPI 2 - Severe Delay Rate (%)", "SUM(SevereDelayFlag) / SUM(TripCount) x 100", "<= 10%",
         "Green <= 10% | Amber > 10% to 15% | Red > 15%", "Month-over-Month change (percentage points)", f"{k2:.2f}%", s2, "SevereDelayFlag, TripCount"),
        ("KPI 3 - Average Excess Delay per Trip (min)", "SUM(ExcessDelayMinutes) / SUM(TripCount)", "<= 10 min",
         "Green <= 10 | Amber > 10 to 20 | Red > 20", "Month-over-Month change (minutes)", f"{k3:.2f} min", s3, "ExcessDelayMinutes, TripCount"),
        ("KPI 4 - Weekend Delay Uplift (%)", "(AVG delay on weekends / AVG delay on weekdays - 1) x 100, where AVG = SUM(DelayMinutes) / SUM(TripCount), split by Dim_Date.IsWeekend", "<= 5%",
         "Green <= 5% | Amber > 5% to 15% | Red > 15%", "Month-over-Month change (percentage points)", f"{k4:.2f}%", s4, "DelayMinutes, TripCount (+ Dim_Date.IsWeekend)"),
        ("KPI 5 - Peak-Hour Delay Surge (%)", "(AVG delay in peak hours / AVG delay off-peak - 1) x 100, split by Dim_Time.IsPeakHour (known scheduled time only)", "<= 10%",
         "Green <= 10% | Amber > 10% to 25% | Red > 25%", "Month-over-Month change (percentage points)", f"{k5:.2f}%", s5, "DelayMinutes, TripCount (+ Dim_Time.IsPeakHour)"),
    ]
    for name, formula, goal, bands, trend, actual, st, meas in kpis:
        H(name, 2)
        T(["Property", "Definition"], [
            ["Value (formula)", formula], ["Goal", goal], ["Status bands", bands], ["Trend", trend],
            ["Fact measures used", meas], ["Actual value (full period)", actual], ["Current status", st],
        ], [1.8, 5.0], font=8)

    H("6.1 Monthly KPI trend (complete months, from the warehouse)", 2)
    rows = []
    for r in mon.itertuples():
        def f(v, suf=''):
            return '-' if pd.isna(v) else f"{float(v):.2f}{suf}"
        rows.append([f"{int(r.Year)}-{int(r.Month):02d}", f"{int(r.TotalTrips):,}",
                     f(r.PunctualityRatePct), f(r.PunctualityMoMChangePts), f(r.SevereDelayRatePct),
                     f(r.SevereDelayMoMChangePts), f(r.AvgExcessDelayMin), f(r.WeekendUpliftPct), f(r.PeakSurgePct)])
    T(["Month", "Trips", "K1 %", "K1 MoM", "K2 %", "K2 MoM", "K3 min", "K4 %", "K5 %"], rows, font=7)

    H("7. Traceability Matrix (KPI -> fact measure -> physical column)")
    T(["KPI", "Measure / attribute in formula", "Physical location"], [
        ["KPI 1, 2, 3", "TripCount", "Fact_TrainDelay.TripCount"],
        ["KPI 1", "PunctualFlag", "Fact_TrainDelay.PunctualFlag"],
        ["KPI 2", "SevereDelayFlag", "Fact_TrainDelay.SevereDelayFlag"],
        ["KPI 3", "ExcessDelayMinutes", "Fact_TrainDelay.ExcessDelayMinutes"],
        ["KPI 4, 5", "DelayMinutes", "Fact_TrainDelay.DelayMinutes"],
        ["KPI 4", "IsWeekend (slicer)", "Dim_Date.IsWeekend via Fact_TrainDelay.DateSK"],
        ["KPI 5", "IsPeakHour (slicer)", "Dim_Time.IsPeakHour via Fact_TrainDelay.TimeSK"],
        ["Trend (all)", "Year, Month", "Dim_Date.Year, Dim_Date.Month via Fact_TrainDelay.DateSK"],
    ], [1.2, 2.6, 3.0], font=8)
    P("Every measure used in a KPI formula is a physical column of the fact table; every slicing attribute is a physical column of a dimension reachable by a foreign key.")

    H("8. Data Quality Rules & Assumptions")
    B([
        "Source rows with an empty delay value (no measurement, ~1.9 million) are excluded from Fact_TrainDelay instead of being counted as on-time.",
        "Delays whose absolute value exceeds 24 hours (1,440 minutes) are treated as data errors and excluded.",
        "Train numbers are zero-padded to 5 digits so delay/schedule keys match Dim_Train; rows whose station or train is missing from the dimensions are skipped.",
        "Punctual = delay <= 15 minutes (common railway tolerance); Severe = delay > 45 minutes; Excess delay = minutes beyond the 15-minute tolerance.",
        "Season: Winter (Dec-Feb), Summer (Mar-May), Monsoon (Jun-Sep), Post-Monsoon (Oct-Nov). Holidays are a representative list of Indian public holidays.",
        "TimeSK is the scheduled arrival (else departure) time of the stop, found by joining (train, stop sequence) with the timetable; unmatched rows use TimeSK 9999 (Unknown).",
        "The data covers about 13 months, so Year-over-Year trends are not possible; Month-over-Month is used for all KPIs.",
        "All 8 business events are physically loaded: Events 1-4 from the high-frequency telemetry & schedule datasets, Events 5-7 from the operational journey feed, and Event 8 aggregated as a daily periodic snapshot.",
        "The KPI goals are management targets chosen for this assignment; the actual values show the current gap to target.",
    ])

    H("9. Implementation & Reproduction")
    B([
        "sql/01-03: database, dimensional model DDL (8 dimensions, 8 fact tables), indexes and KPI views.",
        "etl/load_dimensions.py: Date/Time generation, Station zone enrichment, SCD2 Train load.",
        "etl/load_fact.py: Fact_TrainSchedule and chunked Fact_TrainDelay load with derived measures, departures, and arrivals.",
        "etl/load_event_facts.py: Fact_JourneyCancellation, Fact_DelayIncident, Fact_WeatherObservation, and Fact_TrainDailySnapshot load.",
        "etl/validate_load.py: 70+ automated checks verifying source reconciliation, RI, SCD2, grain, and checksums.",
        "docker compose up -d  ->  python3 etl/run_etl.py  (CSV files in data/raw/ or the project root).",
    ])
    P(f"Complete Load Result across all 8 Fact Tables: Fact_TrainDelay {n['Fact_TrainDelay']:,} rows; "
      f"Fact_TrainSchedule {n['Fact_TrainSchedule']:,} rows; "
      f"Fact_JourneyDeparture {n['Fact_JourneyDeparture']:,} rows; "
      f"Fact_JourneyArrival {n['Fact_JourneyArrival']:,} rows; "
      f"Fact_JourneyCancellation {n['Fact_JourneyCancellation']:,} rows; "
      f"Fact_DelayIncident {n['Fact_DelayIncident']:,} rows; "
      f"Fact_WeatherObservation {n['Fact_WeatherObservation']:,} rows; "
      f"Fact_TrainDailySnapshot {n['Fact_TrainDailySnapshot']:,} rows.")


if __name__ == "__main__":
    engine = get_engine()
    data = collect(engine)
    build(data)
    doc.save(OUT_DOCX)
    with open(OUT_MD, "w", encoding="utf-8") as fh:
        fh.write("\n".join(md))
    print("Wrote", OUT_DOCX)
    print("Wrote", OUT_MD)
